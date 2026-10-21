"""Exportador de transcrições judiciais revisadas para arquivos Microsoft Word (.docx) profissionais."""

from __future__ import annotations

import io
import logging
import re
import zipfile
from pathlib import Path
from typing import Any

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

logger = logging.getLogger(__name__)

# Paleta de Cores Institucionais STJ / Gabinete GMHB
COR_NAVY_STJ = RGBColor(10, 37, 64)         # #0A2540
COR_AZUL_MEDIO = RGBColor(30, 58, 138)       # #1E3A8A
COR_TEXTO_ESCURO = RGBColor(31, 41, 55)      # #1F2937
COR_TEXTO_MUTED = RGBColor(107, 114, 128)    # #6B7280
COR_LINK_YOUTUBE = RGBColor(29, 78, 216)     # #1D4ED8

HEX_SHADING_HERMAN = "EFF6FF"               # Azul muito suave para destaque do Min. Herman
HEX_BORDER_HERMAN = "1E3A8A"                # Azul marinho STJ
HEX_SHADING_PROCESSO = "F1F5F9"             # Cinza azulado muito claro para cabeçalhos de processo
HEX_BORDER_PROCESSO = "0284C7"              # Azul ciano para marcar apregoamento

TIMESTAMP_RE = re.compile(r"\[(\d{1,2}:\d{2}:\d{2})\]")
SPEAKER_LINE_RE = re.compile(r"^\*\*\[(\d{1,2}:\d{2}:\d{2})\]\s*([^:*]+):\*\*\s*(.*)$")
HEADING_PROCESSO_RE = re.compile(r"^###\s+\[?(\d{1,2}:\d{2}:\d{2})\]?\s*(.*)$")


def _conteudo_docx(dados: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(dados)) as arquivo_zip:
        return {nome: arquivo_zip.read(nome) for nome in arquivo_zip.namelist()}


def _gravar_docx_se_mudou(caminho: Path, dados: bytes) -> bool:
    """Grava o .docx só se o conteúdo mudou.

    O Word guarda a data de gravação dentro do arquivo; sem essa comparação, cada build
    regravaria o mesmo documento e ele apareceria sempre como modificado.
    """
    if caminho.exists():
        try:
            if _conteudo_docx(caminho.read_bytes()) == _conteudo_docx(dados):
                return False
        except zipfile.BadZipFile:
            pass
    caminho.write_bytes(dados)
    return True


def _timestamp_para_segundos(ts: str) -> int:
    if not ts:
        return 0
    partes = [int(p) for p in ts.strip().split(":") if p.isdigit()]
    if len(partes) == 3:
        return partes[0] * 3600 + partes[1] * 60 + partes[2]
    if len(partes) == 2:
        return partes[0] * 60 + partes[1]
    return 0


def _construir_url_tempo(video_url: str, segundos: int) -> str:
    if not video_url or segundos <= 0:
        return video_url or ""
    sep = "&" if "?" in video_url else "?"
    return f"{video_url}{sep}t={segundos}s"


def _adicionar_hyperlink(paragraph, url: str, texto: str, cor: RGBColor = COR_LINK_YOUTUBE, sublinhado: bool = True):
    """Adiciona um hyperlink executável no documento Word."""
    if not url:
        run = paragraph.add_run(texto)
        run.font.color.rgb = cor
        return run

    part = paragraph.part
    r_id = part.relate_to(
        url,
        docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )

    hyperlink = parse_xml(
        f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>'
    )
    new_run = parse_xml(f'<w:r {nsdecls("w")}/>')
    rPr = parse_xml(f'<w:rPr {nsdecls("w")}/>')

    if cor:
        hex_color = f"{cor[0]:02X}{cor[1]:02X}{cor[2]:02X}"
        c = parse_xml(f'<w:color {nsdecls("w")} w:val="{hex_color}"/>')
        rPr.append(c)
    if sublinhado:
        u = parse_xml(f'<w:u {nsdecls("w")} w:val="single"/>')
        rPr.append(u)

    new_run.append(rPr)
    text_node = parse_xml(f'<w:t {nsdecls("w")}>{texto}</w:t>')
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def _aplicar_sombreado_paragrafo(paragraph, cor_hex: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{cor_hex}"/>')
    pPr.append(shd)


def _aplicar_borda_esquerda_paragrafo(paragraph, cor_hex: str = HEX_BORDER_HERMAN, tamanho_oitavos_pt: str = "24") -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="{tamanho_oitavos_pt}" w:space="10" w:color="{cor_hex}"/></w:pBdr>'
    )
    pPr.append(pBdr)


def _configurar_estilos_base(doc: docx.Document) -> None:
    """Configura margens e fonte padrão do documento Word."""
    for section in doc.sections:
        section.top_margin = Inches(0.98)     # ~2,5 cm
        section.bottom_margin = Inches(0.98)  # ~2,5 cm
        section.left_margin = Inches(0.98)    # ~2,5 cm
        section.right_margin = Inches(0.98)   # ~2,5 cm
        section.page_width = Inches(8.27)     # A4
        section.page_height = Inches(11.69)   # A4

        # Cabeçalho institucional discreto
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hdr.paragraph_format.space_after = Pt(2)
        r_hdr = p_hdr.add_run("SUPERIOR TRIBUNAL DE JUSTIÇA — GABINETE DO MINISTRO HERMAN BENJAMIN")
        r_hdr.font.name = "Calibri"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = COR_TEXTO_MUTED

        # Rodapé com identificação
        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_ftr = p_ftr.add_run("Transcrição Taquigráfica Oficial Saneada — Sessões de Julgamento")
        r_ftr.font.name = "Calibri"
        r_ftr.font.size = Pt(8.5)
        r_ftr.font.color.rgb = COR_TEXTO_MUTED

    # Estilo Normal
    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Calibri"
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = COR_TEXTO_ESCURO
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)


def _adicionar_cabecalho_institucional(
    doc: docx.Document,
    orgao: str,
    data_sessao: str,
    video_url: str,
    duracao: str = "",
    pauta: dict[str, Any] | None = None,
) -> None:
    """Insere o cabeçalho oficial com metadados da sessão."""
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(0)
    p_top.paragraph_format.space_after = Pt(2)
    r1 = p_top.add_run("SUPERIOR TRIBUNAL DE JUSTIÇA")
    r1.font.bold = True
    r1.font.size = Pt(14)
    r1.font.color.rgb = COR_NAVY_STJ

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    r2 = p_sub.add_run("Gabinete do Ministro Herman Benjamin")
    r2.font.bold = True
    r2.font.size = Pt(11.5)
    r2.font.color.rgb = COR_AZUL_MEDIO

    p_doc = doc.add_paragraph()
    p_doc.paragraph_format.space_before = Pt(0)
    p_doc.paragraph_format.space_after = Pt(12)
    r3 = p_doc.add_run("Transcrição Taquigráfica Revisada e Saneada")
    r3.font.size = Pt(12)
    r3.font.color.rgb = RGBColor(55, 65, 81)

    dados: list[tuple[str, str]] = [
        ("Órgão Colegiado:", orgao or "Primeira Turma / Primeira Seção"),
        ("Data da Sessão:", data_sessao or "Data não informada"),
        ("Transmissão Original:", video_url or "Vídeo do STJ no YouTube"),
    ]
    if pauta and pauta.get("url_pauta_pdf"):
        dados.append(("Pauta Oficial (STJ):", str(pauta["url_pauta_pdf"])))
    if pauta and pauta.get("url_ata_pdf"):
        dados.append(("Ata de Julgamento:", str(pauta["url_ata_pdf"])))
    dados.append(
        ("Padrão Editorial:", "Taquigrafia revisada sem ruídos de áudio, com falas individualizadas e minutagem.")
    )

    # Tabela com Metadados
    table = doc.add_table(rows=len(dados), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    larguras = [Inches(1.8), Inches(4.5)]

    for row_idx, (rotulo, valor) in enumerate(dados):
        row = table.rows[row_idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width, cell_val.width = larguras[0], larguras[1]

        # Estilo da célula do rótulo
        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(2)
        p_lbl.paragraph_format.space_before = Pt(2)
        r_l = p_lbl.add_run(rotulo)
        r_l.font.bold = True
        r_l.font.size = Pt(10)
        r_l.font.color.rgb = COR_NAVY_STJ

        # Estilo da célula de valor
        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_after = Pt(2)
        p_val.paragraph_format.space_before = Pt(2)
        if rotulo == "Transmissão Original:" and valor.startswith("http"):
            _adicionar_hyperlink(p_val, valor, "Assistir gravação oficial no YouTube", COR_LINK_YOUTUBE)
        elif rotulo == "Pauta Oficial (STJ):" and valor.startswith("http"):
            total_it = pauta.get("total_itens", 0) if pauta else 0
            _adicionar_hyperlink(
                p_val,
                valor,
                f"Abrir Pauta Oficial em PDF ({total_it} itens na Ordem do Dia)",
                COR_LINK_YOUTUBE,
            )
        elif rotulo == "Ata de Julgamento:" and valor.startswith("http"):
            _adicionar_hyperlink(p_val, valor, "Abrir Ata Oficial de Julgamento em PDF", COR_LINK_YOUTUBE)
        else:
            r_v = p_val.add_run(valor)
            r_v.font.size = Pt(10)
            r_v.font.color.rgb = COR_TEXTO_ESCURO

        # Shading suave nas linhas da tabela
        shd_hex = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c in (cell_lbl, cell_val):
            c_shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{shd_hex}"/>')
            c._tc.get_or_add_tcPr().append(c_shd)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(8)
    p_spacer.paragraph_format.space_after = Pt(8)


def _adicionar_quadro_pauta_docx(doc: docx.Document, pauta: dict[str, Any] | None) -> None:
    """Insere o quadro sinótico da Pauta / Ordem do Dia no início do documento Word."""
    if not pauta or not pauta.get("itens"):
        return

    itens = pauta["itens"]
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(10)
    p_sec.paragraph_format.space_after = Pt(6)
    p_sec.paragraph_format.keep_with_next = True
    r_sec = p_sec.add_run(f"Pauta e Ordem do Dia Oficial ({len(itens)} processos)")
    r_sec.font.bold = True
    r_sec.font.size = Pt(13)
    r_sec.font.color.rgb = COR_NAVY_STJ

    table = doc.add_table(rows=len(itens) + 1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    larguras = [Inches(0.5), Inches(1.75), Inches(1.45), Inches(1.35), Inches(1.25)]
    cabecalhos = ["Item", "Processo", "Relator / Vistor", "Assunto", "Resultado"]

    hdr_row = table.rows[0]
    for col_idx, titulo_col in enumerate(cabecalhos):
        cell = hdr_row.cells[col_idx]
        cell.width = larguras[col_idx]
        p_c = cell.paragraphs[0]
        p_c.paragraph_format.space_before = Pt(2)
        p_c.paragraph_format.space_after = Pt(2)
        r_c = p_c.add_run(titulo_col)
        r_c.font.bold = True
        r_c.font.size = Pt(9)
        r_c.font.color.rgb = RGBColor(255, 255, 255)
        c_shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0A2540"/>')
        cell._tc.get_or_add_tcPr().append(c_shd)

    for idx, it in enumerate(itens, start=1):
        row = table.rows[idx]
        proc_str = it.get("processo") or ""
        if it.get("tema_repetitivo"):
            proc_str = f"{proc_str} ({it['tema_repetitivo']})"
        rel_str = it.get("relator") or ""
        if it.get("vistor"):
            rel_str = f"{rel_str} (Vista: {it['vistor']})"

        valores = [
            f"{int(it.get('item') or idx):02d}",
            proc_str,
            rel_str,
            str(it.get("assunto") or "-"),
            str(it.get("resultado_oficial") or "-"),
        ]
        shd_hex = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(valores):
            cell = row.cells[col_idx]
            cell.width = larguras[col_idx]
            p_c = cell.paragraphs[0]
            p_c.paragraph_format.space_before = Pt(1.5)
            p_c.paragraph_format.space_after = Pt(1.5)
            r_c = p_c.add_run(val)
            r_c.font.size = Pt(8.5)
            if col_idx in (0, 1):
                r_c.font.bold = True
                r_c.font.color.rgb = COR_NAVY_STJ
            else:
                r_c.font.color.rgb = COR_TEXTO_ESCURO
            c_shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{shd_hex}"/>')
            cell._tc.get_or_add_tcPr().append(c_shd)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(6)
    p_sp.paragraph_format.space_after = Pt(6)


def _renderizar_texto_com_negrito_inline(paragraph, texto: str, cor_padrao: RGBColor = COR_TEXTO_ESCURO) -> None:
    """Analisa marcações markdown inline simples como **texto em negrito**."""
    partes = re.split(r"(\*\*.*?\*\*)", texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**") and len(parte) >= 4:
            run = paragraph.add_run(parte[2:-2])
            run.font.bold = True
            run.font.color.rgb = cor_padrao
        else:
            run = paragraph.add_run(parte)
            run.font.color.rgb = cor_padrao


def exportar_transcricao_para_docx(
    markdown_transcricao: str,
    output_path: Path | str,
    video_url: str = "",
    orgao: str = "",
    data_sessao: str = "",
    titulo_sessao: str = "",
    pauta: dict[str, Any] | None = None,
) -> Path:
    """
    Converte o Markdown de uma transcrição saneada do STJ em um documento
    Microsoft Word (.docx) com formatação institucional e tipografia limpa.
    """
    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()
    _configurar_estilos_base(doc)

    # Inserir cabeçalho institucional
    _adicionar_cabecalho_institucional(
        doc=doc,
        orgao=orgao,
        data_sessao=data_sessao,
        video_url=video_url,
        pauta=pauta,
    )

    # Inserir quadro sinótico da Pauta / Ordem do Dia quando disponível
    _adicionar_quadro_pauta_docx(doc=doc, pauta=pauta)

    # Processar o corpo da transcrição
    linhas = (markdown_transcricao or "").replace("\r\n", "\n").replace("\r", "\n").splitlines()

    pulando_cabecalho_inicial = True
    for linha in linhas:
        texto = linha.strip()
        if not texto:
            continue

        # Ignorar o cabeçalho original em Markdown (já renderizado na tabela institucional)
        if pulando_cabecalho_inicial:
            if texto.startswith("# ") or texto.startswith("- **Órgão") or texto.startswith("- **Data") or texto.startswith("- **Duração") or texto.startswith("- **Vídeo") or texto.startswith("- **Processamento") or texto == "---":
                continue
            if texto.startswith("## "):
                pulando_cabecalho_inicial = False
            else:
                continue

        # 1. Título de Seção Principal (## ...)
        if texto.startswith("## "):
            titulo_sec = texto[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(titulo_sec)
            r.font.bold = True
            r.font.size = Pt(13)
            r.font.color.rgb = COR_NAVY_STJ
            continue

        # 2. Subtítulo de Apregoamento / Processo / Bloco (### [HH:MM:SS] ...)
        match_proc = HEADING_PROCESSO_RE.match(texto)
        if match_proc or texto.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            _aplicar_sombreado_paragrafo(p, HEX_SHADING_PROCESSO)
            _aplicar_borda_esquerda_paragrafo(p, HEX_BORDER_PROCESSO, tamanho_oitavos_pt="16")

            if match_proc:
                ts, desc_proc = match_proc.group(1), match_proc.group(2).strip()
                segs = _timestamp_para_segundos(ts)
                url_tempo = _construir_url_tempo(video_url, segs)

                if url_tempo:
                    _adicionar_hyperlink(p, url_tempo, f"[{ts}]", COR_AZUL_MEDIO, sublinhado=False)
                    p.add_run(" ")
                else:
                    r_ts = p.add_run(f"[{ts}] ")
                    r_ts.font.bold = True
                    r_ts.font.color.rgb = COR_AZUL_MEDIO

                r_desc = p.add_run(desc_proc)
                r_desc.font.bold = True
                r_desc.font.color.rgb = COR_NAVY_STJ
            else:
                r = p.add_run(texto[4:].strip())
                r.font.bold = True
                r.font.color.rgb = COR_NAVY_STJ
            continue

        # 3. Intervenção de Orador (**[HH:MM:SS] Orador:** Texto)
        match_speaker = SPEAKER_LINE_RE.match(texto)
        if match_speaker:
            ts, orador, fala = match_speaker.group(1), match_speaker.group(2).strip(), match_speaker.group(3).strip()
            segs = _timestamp_para_segundos(ts)
            url_tempo = _construir_url_tempo(video_url, segs)

            p = doc.add_paragraph()
            eh_herman = "herman benjamin" in orador.lower()

            if eh_herman:
                # Destaque gráfico especial para o Min. Herman Benjamin
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                _aplicar_sombreado_paragrafo(p, HEX_SHADING_HERMAN)
                _aplicar_borda_esquerda_paragrafo(p, HEX_BORDER_HERMAN, tamanho_oitavos_pt="24")

                if url_tempo:
                    _adicionar_hyperlink(p, url_tempo, f"[{ts}]", COR_AZUL_MEDIO, sublinhado=False)
                    p.add_run(" ")
                else:
                    r_ts = p.add_run(f"[{ts}] ")
                    r_ts.font.bold = True
                    r_ts.font.color.rgb = COR_AZUL_MEDIO

                r_ora = p.add_run(f"{orador}: ")
                r_ora.font.bold = True
                r_ora.font.color.rgb = COR_NAVY_STJ

                _renderizar_texto_com_negrito_inline(p, fala, COR_TEXTO_ESCURO)
            else:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(4)

                if url_tempo:
                    _adicionar_hyperlink(p, url_tempo, f"[{ts}]", COR_TEXTO_MUTED, sublinhado=False)
                    p.add_run(" ")
                else:
                    r_ts = p.add_run(f"[{ts}] ")
                    r_ts.font.color.rgb = COR_TEXTO_MUTED

                r_ora = p.add_run(f"{orador}: ")
                r_ora.font.bold = True
                r_ora.font.color.rgb = COR_TEXTO_ESCURO

                _renderizar_texto_com_negrito_inline(p, fala, COR_TEXTO_ESCURO)
            continue

        # 4. Parágrafo simples ou com timestamp isolado ([HH:MM:SS] Texto)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)

        match_ts_start = re.match(r"^\[(\d{1,2}:\d{2}:\d{2})\]\s*(.*)$", texto)
        if match_ts_start:
            ts, resto = match_ts_start.group(1), match_ts_start.group(2)
            segs = _timestamp_para_segundos(ts)
            url_tempo = _construir_url_tempo(video_url, segs)
            if url_tempo:
                _adicionar_hyperlink(p, url_tempo, f"[{ts}]", COR_TEXTO_MUTED, sublinhado=False)
                p.add_run(" ")
            else:
                r_ts = p.add_run(f"[{ts}] ")
                r_ts.font.color.rgb = COR_TEXTO_MUTED
            _renderizar_texto_com_negrito_inline(p, resto, COR_TEXTO_ESCURO)
        else:
            _renderizar_texto_com_negrito_inline(p, texto, COR_TEXTO_ESCURO)

    buffer = io.BytesIO()
    doc.save(buffer)
    if _gravar_docx_se_mudou(Path(output_file), buffer.getvalue()):
        logger.info("Transcrição exportada para Word com sucesso: %s", output_file)
    else:
        logger.info("Transcrição em Word sem mudanças: %s", output_file)
    return output_file
